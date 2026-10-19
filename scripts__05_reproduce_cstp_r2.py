#!/usr/bin/env python3
"""Reproduce the descriptive numerical findings of the CSTP City–Metro case study.

Read-only: never changes the validated master workbook or reported manuscript data.
Outputs are generated from the processed source tables, not copied from figures.

Requirements: Python >= 3.9, openpyxl >= 3.1.
Example:
 python scripts__05_reproduce_cstp_r2.py --check
 python scripts__05_reproduce_cstp_r2.py --master path/to/master.xlsx --output-dir audit_out
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

MASTER_NAME = "data__processed__TomTom_Malaysia_7Cities_City_Metro_2025_Master.xlsx"
CITIES = (
    "George Town", "Kota Bharu", "Kuala Lumpur", "Seberang Perai",
    "Johor Bahru", "Ipoh", "Kajang",
)
SCALES = ("City", "Metro")
INDICATORS = {
    "congestion": ("avg_congestion_percent", 1),
    "tt10": ("avg_travel_time_10km_seconds", 1),
    "speed": ("avg_speed_rush_hour_kmh", -1),
    "distance": ("avg_distance_15min_km", -1),
}
WEIGHTS = {
    "M0": (1 / 4, 1 / 4, 1 / 4, 1 / 4),
    "M1": (0, 1 / 3, 1 / 3, 1 / 3),
    "M2": (1 / 3, 0, 1 / 3, 1 / 3),
    "M3": (1 / 3, 1 / 3, 0, 1 / 3),
    "M4": (1 / 3, 1 / 3, 1 / 3, 0),
    "M5": (1 / 2, 1 / 6, 1 / 6, 1 / 6),
}
# Published, rounded targets are regression checks ONLY; no result is calculated from them.
EXPECTED_PAI = {
    "George Town": -0.032, "Kota Bharu": 0.045,
    "Kuala Lumpur": 0.049, "Seberang Perai": -0.135,
    "Johor Bahru": -0.038, "Ipoh": 0.027, "Kajang": -0.005,
}
EXPECTED_TAS = {
    "George Town": (13.63, 20.17, 19.43),
    "Kota Bharu": (28.68, 26.44, 25.11),
    "Kuala Lumpur": (14.46, 16.56, 16.53),
    "Seberang Perai": (17.99, 4.00, 3.78),
    "Johor Bahru": (15.49, 27.69, 28.73),
    "Ipoh": (15.04, 33.68, 35.47),
    "Kajang": (7.26, 10.11, 10.42),
}
EXPECTED_SCREENING_J = {
    ("congestion", 2): 1 / 3, ("congestion", 3): 1 / 2,
    ("tt10", 2): 1.0, ("tt10", 3): 3 / 4,
    ("speed", 2): 1 / 3, ("speed", 3): 1 / 2,
}
EXPECTED_MONTHLY_SIGNED = {
    "George Town": 5.0833333333, "Kota Bharu": 12.3333333333,
    "Kuala Lumpur": 4.25, "Seberang Perai": 5.4166666667,
    "Johor Bahru": 0.5833333333, "Ipoh": 4.1666666667,
    "Kajang": -0.8333333333,
}


def load_rows(workbook, sheet):
    rows = workbook[sheet].iter_rows(values_only=True)
    headers = next(rows)
    if len(set(headers)) != len(headers):
        raise ValueError(f"Duplicate headers in {sheet}")
    return [dict(zip(headers, row)) for row in rows if any(x is not None for x in row)]


def numerical(x, label):
    if isinstance(x, bool) or not isinstance(x, (int, float)) or not math.isfinite(x):
        raise ValueError(f"Missing/non-numeric {label}: {x!r}")
    return float(x)


def rows_to_csv(path, records, fields):
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(records)


def paired_records(rows, keys, scale_field="scale", allowed_scales=SCALES):
    groups = defaultdict(dict)
    for r in rows:
        key = tuple(r[name] for name in keys)
        scale = r[scale_field]
        if scale not in allowed_scales:
            raise ValueError(f"Unknown reporting unit: {scale}")
        if scale in groups[key]:
            raise ValueError(f"Duplicate paired cell: {key} | {scale}")
        groups[key][scale] = r
    for key, cell in groups.items():
        if set(cell) != set(allowed_scales):
            raise ValueError(f"Incomplete City–Metro pair: {key}: {list(cell)}")
    return groups


def competition_ranks(scores, higher_worse=True):
    ordered = sorted(scores, key=lambda c: (scores[c] if higher_worse else -scores[c]), reverse=True)
    result = {}
    previous = None
    previous_rank = 0
    for pos, city in enumerate(ordered, 1):
        value = scores[city]
        rank = previous_rank if previous is not None and value == previous else pos
        result[city] = rank
        previous, previous_rank = value, rank
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--master", type=Path, default=Path(__file__).resolve().parent / MASTER_NAME)
    parser.add_argument("--output-dir", type=Path, default=Path("cstp_r2_reproduction_output"))
    parser.add_argument("--check", action="store_true", help="Fail if published rounded results do not match")
    parser.add_argument("--data-dir", type=Path, help="Use de-identified input CSV files instead of the XLSX master")
    parser.add_argument("--build-review-package", type=Path,
                        help="Write a self-contained double-blind reviewer ZIP from the validated master and calculated results")
    args = parser.parse_args()
    if args.data_dir and args.build_review_package:
        parser.error("--build-review-package requires --master, not --data-dir")
    if args.data_dir:
        input_dir = args.data_dir
        stem_map = {
            "annual": "01_annual_overview.csv",
            "monthly": "02_monthly_congestion.csv",
            "rush": "03_rush_hour_clean.csv",
            "heatmap": "04_heatmap_long.csv",
            "source_audit": "09_source_audit.csv",
        }
        sources = {}
        for key, name in stem_map.items():
            f = input_dir / name
            if not f.is_file():
                parser.error(f"Missing anonymous input: {f}")
            with f.open("r", encoding="utf-8-sig", newline="") as inp:
                sources[key] = list(csv.DictReader(inp))
        # Numeric coercion for the public-facing, self-contained CSV reproduction route.
        numeric_cols = {
            "year", "month_number", "hour_24h", "tt10_seconds",
            "avg_congestion_percent", "avg_travel_time_10km_seconds",
            "avg_speed_rush_hour_kmh", "avg_distance_15min_km",
            "congestion_percent", "value_numeric",
        }
        for rows in sources.values():
            for row in rows:
                for col in numeric_cols & row.keys():
                    if row[col] is not None and row[col].strip() != "":
                        row[col] = float(row[col])
        annual, monthly, rush = sources["annual"], sources["monthly"], sources["rush"]
        heat, source_audit = sources["heatmap"], sources["source_audit"]
        input_digests = {k: hashlib.sha256((input_dir / f).read_bytes()).hexdigest()
                         for k, f in stem_map.items()}
        source_mode = "anonymous_csv"
    else:
        if not args.master.is_file():
            parser.error(f"Validated workbook not found: {args.master}. Pass --master explicitly.")
        wb = load_workbook(args.master, read_only=True, data_only=True)
        required_sheets = {
            "01_annual_overview", "02_monthly_congestion", "03_rush_hour_clean",
            "04_heatmap_long", "09_source_audit", "11_validation",
        }
        if not required_sheets.issubset(wb.sheetnames):
            raise ValueError(f"Missing master sheets: {required_sheets - set(wb.sheetnames)}")
        annual = load_rows(wb, "01_annual_overview")
        monthly = load_rows(wb, "02_monthly_congestion")
        rush = load_rows(wb, "03_rush_hour_clean")
        heat = load_rows(wb, "04_heatmap_long")
        source_audit = load_rows(wb, "09_source_audit")
        wb.close()
        input_digests = {"validated_master_workbook": hashlib.sha256(args.master.read_bytes()).hexdigest()}
        source_mode = "validated_master_xlsx"
    expected_counts = {"annual": 14, "monthly": 336, "rush": 28, "heatmap": 7056, "source_audit": 70}
    actual_counts = {"annual": len(annual), "monthly": len(monthly), "rush": len(rush),
                     "heatmap": len(heat), "source_audit": len(source_audit)}
    for key, n in expected_counts.items():
        if actual_counts[key] != n:
            raise ValueError(f"{key}: expected {n} rows, found {actual_counts[key]}")

    output = args.output_dir
    output.mkdir(parents=True, exist_ok=True)
    errors = []

    # RQ1: four annual indicator comparisons, normalized signed and burden-adjusted gaps.
    by_year = paired_records([r for r in annual if r["year"] == 2025], ("city", "year"))
    if len(by_year) != 7:
        raise ValueError(f"Expected 7 annual case pairs; found {len(by_year)}")
    annual_out, abs_by_city, level_by_unit = [], {}, {i: {} for i in ("congestion", "tt10", "speed")}
    for (city, year), pair in by_year.items():
        abs_by_city[city] = {}
        for ind, (column, burden_sign) in INDICATORS.items():
            cv = numerical(pair["City"][column], f"{city} City {column}")
            mv = numerical(pair["Metro"][column], f"{city} Metro {column}")
            denom = (cv + mv) / 2
            if denom <= 0:
                raise ValueError(f"Non-positive paired mean: {city} {ind}")
            signed = 100 * (cv - mv) / denom
            adjusted = burden_sign * signed
            abs_by_city[city][ind] = abs(adjusted)
            if ind in level_by_unit:
                level_by_unit[ind][city] = {"City": cv, "Metro": mv}
            annual_out.append(dict(city=city, indicator=ind, city_value=cv,
                                   metro_value=mv, signed_nd_percent=signed,
                                   burden_nd_percent=adjusted, abs_sensitivity_percent=abs(adjusted)))
    rows_to_csv(output / "annual_sensitivity.csv", annual_out,
                ["city", "indicator", "city_value", "metro_value", "signed_nd_percent",
                 "burden_nd_percent", "abs_sensitivity_percent"])
    annual_means = {ind: sum(abs_by_city[c][ind] for c in CITIES) / 7 for ind in INDICATORS}
    if args.check:
        for ind, published in {"congestion": 11.4, "tt10": 19.2, "speed": 18.3, "distance": 19.0}.items():
            if abs(annual_means[ind] - published) > 0.051:
                errors.append(f"Annual {ind}: {annual_means[ind]:.4f} vs published {published:.1f}")

    # RQ1: M0–M5, descriptive case sensitivity, no model fitting or new experiment.
    weight_out, model_scores, rank_models = [], {}, {}
    for model, weights in WEIGHTS.items():
        sc = {city: sum(weights[j] * abs_by_city[city][ind]
                        for j, ind in enumerate(INDICATORS)) for city in CITIES}
        model_scores[model] = sc
        ranks = competition_ranks(sc)
        rank_models[model] = ranks
        for city in CITIES:
            weight_out.append(dict(model=model, city=city, score_percent=sc[city],
                                   rank=ranks[city], weights=";".join(f"{x:.9f}" for x in weights)))
    rows_to_csv(output / "robustness_M0_M5.csv", weight_out,
                ["model", "city", "score_percent", "rank", "weights"])
    top2 = {model: sorted([city for city, rank in ranks.items() if rank <= 2])
            for model, ranks in rank_models.items()}
    if args.check:
        for model, group in top2.items():
            if group != sorted(("Kota Bharu", "Ipoh")):
                errors.append(f"{model} Top-2 = {group}, expected Kota Bharu and Ipoh")

    # RQ2: exact published competition ranking / tie-inclusive k=2,3 screening.
    screening_out = []
    for ind, city_vals in level_by_unit.items():
        field_direction = INDICATORS[ind][1]
        rc = competition_ranks({c: v["City"] for c, v in city_vals.items()}, higher_worse=field_direction == 1)
        rm = competition_ranks({c: v["Metro"] for c, v in city_vals.items()}, higher_worse=field_direction == 1)
        for k in (2, 3):
            sc = {c for c, rank in rc.items() if rank <= k}
            sm = {c for c, rank in rm.items() if rank <= k}
            jaccard = len(sc & sm) / len(sc | sm)
            screening_out.append(dict(indicator=ind, top_k=k, city_set="; ".join(sorted(sc)),
                                      metro_set="; ".join(sorted(sm)),
                                      shared_members=len(sc & sm), union_members=len(sc | sm),
                                      switched_members=len(sc ^ sm), jaccard=jaccard,
                                      identical_sets=(sc == sm)))
            if args.check and abs(jaccard - EXPECTED_SCREENING_J[(ind, k)]) > 1e-9:
                errors.append(f"Screening {ind} Top-{k}: J={jaccard} vs published {EXPECTED_SCREENING_J[(ind,k)]}")
    rows_to_csv(output / "screening_top2_top3.csv", screening_out,
                ["indicator", "top_k", "city_set", "metro_set", "shared_members",
                 "union_members", "switched_members", "jaccard", "identical_sets"])
    if args.check and sum(not row["identical_sets"] for row in screening_out) != 5:
        errors.append("Expected 5/6 screening scenarios with different shortlists")

    # RQ3: Equation (9), 28 clean AM/PM TT10 observations (seconds).
    rp = paired_records([r for r in rush if r["year"] == 2025], ("city", "year", "period"))
    rush_by_city = defaultdict(dict)
    for (city, year, period), pair in rp.items():
        period = str(period).upper().strip()
        if period not in ("AM", "PM"):
            raise ValueError(f"Unexpected rush-hour period: {period}")
        for scale in SCALES:
            t = numerical(pair[scale]["tt10_seconds"], f"rush {city}|{scale}|{period}")
            rush_by_city[(city, scale)][period] = t
    pai_out = []
    for city in CITIES:
        p = {}
        for scale in SCALES:
            slots = rush_by_city[(city, scale)]
            if set(slots) != {"AM", "PM"}:
                raise ValueError(f"Incomplete AM/PM record: {city} {scale}")
            p[scale] = 2 * (slots["PM"] - slots["AM"]) / (slots["PM"] + slots["AM"])
        delta = p["City"] - p["Metro"]
        pai_out.append(dict(city=city, city_AM_seconds=rush_by_city[(city, "City")]["AM"],
                            city_PM_seconds=rush_by_city[(city, "City")]["PM"],
                            metro_AM_seconds=rush_by_city[(city, "Metro")]["AM"],
                            metro_PM_seconds=rush_by_city[(city, "Metro")]["PM"],
                            city_PAI=p["City"], metro_PAI=p["Metro"], delta_PAI=delta))
        if args.check and abs(delta - EXPECTED_PAI[city]) > 0.000501:
            errors.append(f"PAI {city}: {delta:.6f} vs published {EXPECTED_PAI[city]:.3f}")
    rows_to_csv(output / "peak_asymmetry_PAI.csv", pai_out,
                ["city", "city_AM_seconds", "city_PM_seconds", "metro_AM_seconds", "metro_PM_seconds",
                 "city_PAI", "metro_PAI", "delta_PAI"])

    # RQ3: Equation (7), all 7*3*168 pairs; ratio-of-sums, NOT a mean of slotwise percentages.
    hp = paired_records([r for r in heat if r["year"] == 2025],
                        ("city", "year", "indicator", "day_of_week", "hour_24h"))
    totals = defaultdict(lambda: {"numerator": 0.0, "denominator": 0.0, "cells": 0})
    for (city, year, indicator, day, hour), pair in hp.items():
        cv = numerical(pair["City"]["value_numeric"], f"heat {city}|{indicator}|City")
        mv = numerical(pair["Metro"]["value_numeric"], f"heat {city}|{indicator}|Metro")
        rec = totals[(city, indicator)]
        rec["numerator"] += abs(cv - mv)
        rec["denominator"] += (cv + mv) / 2
        rec["cells"] += 1
    tas_out = []
    expected_indicator_order = ("congestion", "travel_time", "speed")
    for city in CITIES:
        for idx, indicator in enumerate(expected_indicator_order):
            s = totals[(city, indicator)]
            if s["cells"] != 168 or s["denominator"] <= 0:
                raise ValueError(f"Invalid TAS slots: {city} {indicator}: {s}")
            tas = 100 * s["numerator"] / s["denominator"]
            tas_out.append(dict(city=city, indicator=indicator, paired_slots=168,
                                sum_absolute_difference=s["numerator"], sum_paired_mean=s["denominator"],
                                TAS_percent=tas))
            if args.check and abs(tas - EXPECTED_TAS[city][idx]) > 0.00501:
                errors.append(f"TAS {city}/{indicator}: {tas:.6f} vs published {EXPECTED_TAS[city][idx]:.2f}")
    rows_to_csv(output / "temporal_absolute_sensitivity_TAS.csv", tas_out,
                ["city", "indicator", "paired_slots", "sum_absolute_difference", "sum_paired_mean", "TAS_percent"])

    # RQ3: Equation (8), 2025-only 12 paired monthly congestion values.
    mp = paired_records([r for r in monthly if r["year"] == 2025],
                        ("city", "year", "month_number"))
    mtotals = defaultdict(lambda: {"numerator": 0.0, "denominator": 0.0, "signed": 0.0, "months": 0})
    for (city, year, month), pair in mp.items():
        cv = numerical(pair["City"]["congestion_percent"], f"monthly {city} City {month}")
        mv = numerical(pair["Metro"]["congestion_percent"], f"monthly {city} Metro {month}")
        s = mtotals[city]
        s["numerator"] += abs(cv - mv)
        s["denominator"] += (cv + mv) / 2
        s["signed"] += cv - mv
        s["months"] += 1
    monthly_out = []
    for city in CITIES:
        s = mtotals[city]
        if s["months"] != 12 or s["denominator"] <= 0:
            raise ValueError(f"Invalid monthly coverage: {city}: {s}")
        monthly_out.append(dict(city=city, paired_months=12,
                                TAS_monthly_percent=100 * s["numerator"] / s["denominator"],
                                mean_signed_gap_percentage_points=s["signed"] / 12))
        if args.check and abs(s["signed"] / 12 - EXPECTED_MONTHLY_SIGNED[city]) > 1e-6:
            errors.append(f"Monthly mean signed difference: {city}")
    rows_to_csv(output / "monthly_congestion_2025.csv", monthly_out,
                ["city", "paired_months", "TAS_monthly_percent", "mean_signed_gap_percentage_points"])

    # Preserve unrounded values when comparing TAS and annual sensitivity.
    johor_tas = next(x["TAS_percent"] for x in tas_out if x["city"] == "Johor Bahru" and x["indicator"] == "congestion")
    johor_delta = johor_tas - abs_by_city["Johor Bahru"]["congestion"]
    if args.check and abs(johor_delta - 13.71) > 0.00501:
        errors.append(f"Johor Bahru rounded TAS-minus-annual contrast: {johor_delta}")

    summary = {
        "script": "scripts__05_reproduce_cstp_r2.py",
        "source_mode": source_mode,
        "input_sha256": input_digests,
        "input_rows": actual_counts,
        "paired_heatmap_cells": len(hp),
        "annual_mean_absolute_sensitivity_percent": annual_means,
        "composite_top2_by_model": top2,
        "screening_nonidentical_out_of_6": sum(not x["identical_sets"] for x in screening_out),
        "PAI_cases_checked": len(pai_out),
        "TAS_city_indicator_cases_checked": len(tas_out),
        "monthly_cases_checked": len(monthly_out),
        "johor_bahru_TAS_minus_annual_congestion_pp_unrounded": johor_delta,
        "check_requested": args.check,
        "checks_passed": (not errors) if args.check else None,
        "errors": errors,
        "method_notes": [
            "The script computes descriptive comparisons of provider-defined records only.",
            "TAS is a ratio of summed absolute paired differences to summed paired means.",
            "Seven independent urban cases; 168 time slots per profile are not independent cities.",
            "These analyses do not validate either reporting unit against external observations.",
        ],
    }
    (output / "reproduction_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    if errors:
        raise SystemExit("FAIL: " + " | ".join(errors))
    print("PASS: Reproduction completed; original input data were not modified."
          + (" Rounded result checks passed." if args.check else " Rounded result checks were not requested."))
    if args.build_review_package:
        # Only purpose-specific processed indicator columns and de-identified audit fields.
        safe_fields = {
            "annual": ["city", "year", "scale", "avg_congestion_percent",
                       "avg_travel_time_10km_seconds", "avg_speed_rush_hour_kmh",
                       "avg_distance_15min_km"],
            "monthly": ["city", "year", "scale", "month_number", "congestion_percent"],
            "rush": ["city", "year", "scale", "period", "tt10_seconds"],
            "heatmap": ["city", "year", "scale", "indicator", "day_of_week",
                        "hour_24h", "value_numeric"],
        }
        base = output / "_review_package_staging"
        data = base / "data"
        data.mkdir(parents=True, exist_ok=True)
        for key, filename in [("annual", "01_annual_overview.csv"),
                              ("monthly", "02_monthly_congestion.csv"),
                              ("rush", "03_rush_hour_clean.csv"),
                              ("heatmap", "04_heatmap_long.csv")]:
            values = {"annual": annual, "monthly": monthly, "rush": rush, "heatmap": heat}[key]
            fields = safe_fields[key]
            if any(not set(fields).issubset(x) for x in values):
                raise ValueError(f"Missing required fields in anonymous {key} export")
            rows_to_csv(data / filename, [{k: r[k] for k in fields} for r in values], fields)
        # The 70 audit records are kept as non-identifying documentation;
        # no author names, user paths, provider screenshots or access credentials are exported.
        audit_allowed = ["city", "scale", "indicator_module", "source_platform",
                         "source_url", "screenshot_file", "source_image",
                         "capture_date", "extraction_method"]
        audit_fields = [c for c in audit_allowed if source_audit and c in source_audit[0]]
        if not audit_fields:
            # Minimal structural fixture support; a production audit should have
            # explanatory source fields. Never fabricate missing provenance.
            audit_fields = [k for k in ("record_id",) if source_audit and k in source_audit[0]]
        if not audit_fields:
            raise ValueError("Audit table has no shareable provenance columns")
        rows_to_csv(data / "09_source_audit.csv",
                    [{k: r[k] for k in audit_fields} for r in source_audit], audit_fields)
        shutil.copy2(Path(__file__), base / "reproduce_review_results.py")
        for name in ["annual_sensitivity.csv", "robustness_M0_M5.csv",
                     "screening_top2_top3.csv", "peak_asymmetry_PAI.csv",
                     "temporal_absolute_sensitivity_TAS.csv", "monthly_congestion_2025.csv",
                     "reproduction_summary.json"]:
            shutil.copy2(output / name, base / name)
        (base / "README_REVIEWERS.txt").write_text(
            "Anonymous processed-data verification package for a seven-city City–Metro case study.\n"
            "Requires: Python 3.9+ and openpyxl. Install: pip install openpyxl\n"
            "Execute from this extracted folder:\n"
            "  python reproduce_review_results.py --data-dir data --check --output-dir rerun\n"
            "The main processed indicators (annual, monthly, rush hour, hourly profiles),\n"
            "a de-identified source-audit index, and derived outputs are provided.\n"
            "Inputs are derived from provider-reported records; not independent validation.\n"
            "Third-party screenshots, proprietary geometries and algorithms are excluded.\n"
            "Seven cities are the cross-case sample; the 168 slots are profiles, not\n"
            "independent cross-city observations.\n", encoding="utf-8")
        target = args.build_review_package.resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as outzip:
            for f in sorted(base.rglob("*")):
                if f.is_file():
                    outzip.write(f, arcname=f.relative_to(base).as_posix())
        print(f"CREATED reviewer ZIP: {target}")
        # Keep normal reproduction outputs, remove intermediate data export after bundling.
        shutil.rmtree(base)


if __name__ == "__main__":
    main()
